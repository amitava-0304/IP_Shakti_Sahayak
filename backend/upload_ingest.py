import os
import re
import uuid
import gc
from pathlib import Path

import chromadb
from dotenv import load_dotenv
from pypdf import PdfReader
from docx import Document

import fitz  # PyMuPDF
import pytesseract
from PIL import Image


# =========================================================
# PATHS / ENVIRONMENT
# =========================================================

PROJECT_ROOT = Path(__file__).resolve().parent.parent

load_dotenv(
    PROJECT_ROOT / ".env"
)

STORAGE_ROOT = Path(
    os.getenv(
        "STORAGE_ROOT",
        str(PROJECT_ROOT)
    )
)

VECTOR_FOLDER = (
    STORAGE_ROOT / "chroma_db"
)

UPLOADED_TEXT_FOLDER = (
    STORAGE_ROOT / "uploaded_text"
)

VECTOR_FOLDER.mkdir(
    parents=True,
    exist_ok=True
)

UPLOADED_TEXT_FOLDER.mkdir(
    parents=True,
    exist_ok=True
)

UPLOAD_COLLECTION = (
    "ip_sakti_uploads"
)

client = chromadb.PersistentClient(
    path=str(VECTOR_FOLDER)
)


def get_upload_collection():
    return client.get_or_create_collection(
        name=UPLOAD_COLLECTION
    )


# =========================================================
# LOW-MEMORY SETTINGS
# =========================================================

OCR_LANGUAGES = os.getenv(
    "OCR_LANGUAGES",
    "eng+ben+hin"
)

OCR_DPI = int(
    os.getenv(
        "OCR_DPI",
        "110"
    )
)

OCR_MIN_PAGE_TEXT = int(
    os.getenv(
        "OCR_MIN_PAGE_TEXT",
        "20"
    )
)

UPLOAD_CHUNK_SIZE = int(
    os.getenv(
        "UPLOAD_CHUNK_SIZE",
        "900"
    )
)

UPLOAD_CHUNK_OVERLAP = int(
    os.getenv(
        "UPLOAD_CHUNK_OVERLAP",
        "120"
    )
)

CHROMA_BATCH_SIZE = int(
    os.getenv(
        "CHROMA_BATCH_SIZE",
        "25"
    )
)


# =========================================================
# TEXT HELPERS
# =========================================================

def clean_text(text):
    text = text or ""

    text = text.replace(
        "\x00",
        " "
    )

    text = re.sub(
        r"[ \t]+",
        " ",
        text
    )

    text = re.sub(
        r"\n{3,}",
        "\n\n",
        text
    )

    return text.strip()


def chunk_text(
    text,
    chunk_size=UPLOAD_CHUNK_SIZE,
    overlap=UPLOAD_CHUNK_OVERLAP
):
    text = clean_text(text)

    if not text:
        return []

    chunks = []

    start = 0

    text_length = len(text)

    while start < text_length:

        end = min(
            start + chunk_size,
            text_length
        )

        chunk = (
            text[start:end]
            .strip()
        )

        if chunk:
            chunks.append(chunk)

        if end >= text_length:
            break

        start = max(
            end - overlap,
            start + 1
        )

    return chunks


# =========================================================
# OCR ONE PAGE
# =========================================================

def ocr_pdf_page(
    pdf_document,
    page_index
):
    page = None
    pix = None
    image = None

    try:

        page = pdf_document.load_page(
            page_index
        )

        pix = page.get_pixmap(
            dpi=OCR_DPI,
            alpha=False
        )

        image = Image.frombytes(
            "RGB",
            [
                pix.width,
                pix.height
            ],
            pix.samples
        )

        text = (
            pytesseract.image_to_string(
                image,
                lang=OCR_LANGUAGES,
                config="--psm 6"
            )
        )

        return clean_text(
            text
        )

    finally:

        if image is not None:
            image.close()

        del image
        del pix
        del page

        gc.collect()


# =========================================================
# PDF EXTRACTION
# =========================================================

def extract_pdf_sections(
    file_path
):
    """
    Low-memory hybrid PDF extraction.

    For each page:
    1. Try pypdf.
    2. If page has little/no text, OCR only that page.
    3. Release image memory before moving to next page.
    """

    reader = PdfReader(
        str(file_path)
    )

    pdf_document = fitz.open(
        str(file_path)
    )

    sections = []

    ocr_pages = 0

    try:

        total_pages = len(
            reader.pages
        )

        for index in range(
            total_pages
        ):

            page_number = (
                index + 1
            )

            text = ""

            try:

                text = (
                    reader.pages[index]
                    .extract_text()
                    or ""
                )

                text = clean_text(
                    text
                )

            except Exception as error:

                print(
                    f"pypdf extraction failed "
                    f"on page {page_number}: "
                    f"{error}",
                    flush=True
                )

                text = ""

            used_ocr = False

            if (
                len(text)
                < OCR_MIN_PAGE_TEXT
            ):

                print(
                    f"OCR page {page_number}...",
                    flush=True
                )

                try:

                    text = ocr_pdf_page(
                        pdf_document,
                        index
                    )

                    if text:
                        used_ocr = True
                        ocr_pages += 1

                except Exception as error:

                    print(
                        f"OCR failed on page "
                        f"{page_number}: "
                        f"{error}",
                        flush=True
                    )

                    text = ""

            if text:

                sections.append(
                    {
                        "page":
                            page_number,

                        "text":
                            text,

                        "ocr":
                            used_ocr
                    }
                )

            gc.collect()

    finally:

        pdf_document.close()

        del reader
        del pdf_document

        gc.collect()

    return (
        sections,
        ocr_pages
    )


# =========================================================
# TXT EXTRACTION
# =========================================================

def extract_txt(
    file_path
):
    with open(
        file_path,
        "r",
        encoding="utf-8",
        errors="ignore"
    ) as file:

        return clean_text(
            file.read()
        )


# =========================================================
# DOCX EXTRACTION
# =========================================================

def extract_docx(
    file_path
):
    document = Document(
        str(file_path)
    )

    parts = []

    for paragraph in (
        document.paragraphs
    ):

        text = clean_text(
            paragraph.text
        )

        if text:
            parts.append(
                text
            )

    for table in (
        document.tables
    ):

        for row in table.rows:

            cells = [
                clean_text(
                    cell.text
                )
                for cell in row.cells
            ]

            line = " | ".join(
                cell
                for cell in cells
                if cell
            )

            if line:
                parts.append(
                    line
                )

    return clean_text(
        "\n".join(parts)
    )


# =========================================================
# SAVE EXTRACTED TEXT
# =========================================================

def save_extracted_text(
    source_filename,
    sections
):
    output_name = (
        Path(source_filename).stem
        + ".txt"
    )

    output_path = (
        UPLOADED_TEXT_FOLDER
        / output_name
    )

    with open(
        output_path,
        "w",
        encoding="utf-8"
    ) as file:

        for section in sections:

            page = section.get(
                "page",
                "Unknown"
            )

            if (
                page
                != "Unknown"
            ):

                file.write(
                    f"===== PAGE {page} =====\n"
                )

            file.write(
                section.get(
                    "text",
                    ""
                ).strip()
            )

            file.write(
                "\n\n"
            )

    return output_path


# =========================================================
# DELETE PREVIOUS VERSION
# =========================================================

def delete_previous_file_chunks(
    collection,
    filename
):
    try:

        collection.delete(
            where={
                "source":
                    filename
            }
        )

    except Exception as error:

        print(
            "Previous upload cleanup "
            f"warning: {error}",
            flush=True
        )


# =========================================================
# STREAM CHUNKS TO CHROMA
# =========================================================

def index_sections(
    filename,
    sections
):
    """
    Index chunks in small batches.
    Avoid keeping all document chunks in memory.
    """

    collection = (
        get_upload_collection()
    )

    delete_previous_file_chunks(
        collection,
        filename
    )

    batch_documents = []
    batch_metadatas = []
    batch_ids = []

    total_chunks = 0

    def flush_batch():

        nonlocal total_chunks

        if not batch_documents:
            return

        collection.add(
            documents=
                batch_documents,

            metadatas=
                batch_metadatas,

            ids=
                batch_ids
        )

        total_chunks += len(
            batch_documents
        )

        batch_documents.clear()
        batch_metadatas.clear()
        batch_ids.clear()

        gc.collect()

    for section in sections:

        page = section.get(
            "page",
            "Unknown"
        )

        ocr_used = bool(
            section.get(
                "ocr",
                False
            )
        )

        chunks = chunk_text(
            section.get(
                "text",
                ""
            )
        )

        for chunk_number, chunk in (
            enumerate(
                chunks,
                start=1
            )
        ):

            batch_documents.append(
                chunk
            )

            batch_metadatas.append(
                {
                    "source":
                        filename,

                    "category":
                        "uploaded",

                    "page":
                        str(page),

                    "uploaded":
                        True,

                    "ocr":
                        ocr_used,

                    "chunk":
                        chunk_number
                }
            )

            batch_ids.append(
                str(
                    uuid.uuid4()
                )
            )

            if (
                len(batch_documents)
                >= CHROMA_BATCH_SIZE
            ):

                flush_batch()

        del chunks

        gc.collect()

    flush_batch()

    return total_chunks


# =========================================================
# MAIN FUNCTION USED BY app.py
# =========================================================

def ingest_uploaded_file(
    file_path
):
    file_path = Path(
        file_path
    )

    if not file_path.exists():

        return {
            "success":
                False,

            "message":
                "Uploaded file was not found."
        }

    extension = (
        file_path.suffix.lower()
    )

    filename = (
        file_path.name
    )

    try:

        ocr_pages = 0

        if (
            extension
            == ".pdf"
        ):

            sections, ocr_pages = (
                extract_pdf_sections(
                    file_path
                )
            )

        elif (
            extension
            == ".txt"
        ):

            text = extract_txt(
                file_path
            )

            sections = [
                {
                    "page":
                        "Unknown",

                    "text":
                        text,

                    "ocr":
                        False
                }
            ]

        elif (
            extension
            == ".docx"
        ):

            text = extract_docx(
                file_path
            )

            sections = [
                {
                    "page":
                        "Unknown",

                    "text":
                        text,

                    "ocr":
                        False
                }
            ]

        else:

            return {
                "success":
                    False,

                "message":
                    "Unsupported file type. "
                    "Please upload PDF, "
                    "TXT or DOCX."
            }

        sections = [
            item
            for item in sections
            if clean_text(
                item.get(
                    "text",
                    ""
                )
            )
        ]

        if not sections:

            return {
                "success":
                    False,

                "message":
                    "No readable text was found. "
                    "OCR was attempted, but "
                    "no usable text could be detected."
            }

        save_extracted_text(
            filename,
            sections
        )

        total_chunks = (
            index_sections(
                filename,
                sections
            )
        )

        del sections

        gc.collect()

        if (
            total_chunks
            <= 0
        ):

            return {
                "success":
                    False,

                "message":
                    "Text was extracted, "
                    "but no searchable chunks "
                    "could be created."
            }

        return {
            "success":
                True,

            "message":
                "Document indexed successfully.",

            "chunks":
                total_chunks,

            "ocr_pages":
                ocr_pages,

            "searchable":
                True
        }

    except Exception as error:

        gc.collect()

        return {
            "success":
                False,

            "message":
                "Document indexing failed: "
                + str(error)
        }
